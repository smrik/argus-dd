import csv

from openpyxl import load_workbook

from argus.layouts import eur_si, write_all
from argus.synth import generate


def test_slovenian_number_format():
    assert eur_si(123456789) == "1.234.567,89"
    assert eur_si(-5) == "-0,05"


def test_messy_layouts_are_written(tmp_path):
    company = generate("trading", 1)
    write_all(company, tmp_path)
    year = company.years[-1].year

    ws = load_workbook(tmp_path / f"tb_{year}_erp.xlsx").active
    total = [r for r in ws.iter_rows(values_only=True) if r[0] == "SKUPAJ"][0]
    assert total[1] == total[2] and total[5] == total[6]  # debits equal credits

    lines = (tmp_path / f"tb_{year}_export.csv").read_text(encoding="cp1250").splitlines()
    assert lines[-1].startswith("Skupaj;;0,00;")

    with open(tmp_path / "tb_tidy.csv", encoding="utf-8") as f:
        assert len(list(csv.DictReader(f))) == sum(len(y.rows) for y in company.years)
