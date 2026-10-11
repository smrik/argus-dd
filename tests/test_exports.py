import csv
import json
import os
import re
import shutil
import subprocess
from collections import defaultdict
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from argus.layouts import write_all, write_erp_xlsx, write_export_csv, write_statements, write_tidy
from argus.aop import AopTree
from argus.synth import generate
from argus.pdf import pdftotext

FIELDS = ("opening", "debit", "credit", "closing")
NUMBER = re.compile(r"-?\d[\d.]*,\d{2}")
FORMATS = ("minimax_pdf", "analitika_pdf", "furs_pdf", "monthly_xlsx", "quirky_xlsx", "tab_csv")


def cents(value):
    return int(Decimal(str(value)) * 100)


def si(value):
    text = value.replace(".", "").replace("\u00a0", "").replace(",", ".")
    negative = text.endswith("-") or text.startswith("(")
    return cents(text.strip("()-")) * (-1 if negative or text.startswith("-") else 1)


def truth(path):
    with path.open(encoding="utf-8", newline="") as f:
        return [(int(r["year"]), r["konto"], r["name"], *(cents(r[k]) for k in FIELDS))
                for r in csv.DictReader(f)]


def pdf_lines(path):
    tool = pdftotext()
    if tool is None:
        pytest.skip("no pdftotext with -table (xpdf 4.x); set PDFTOTEXT")
    return subprocess.run([tool, "-table", str(path), "-"], check=True, capture_output=True,
                          encoding="utf-8").stdout.splitlines()


def read_minimax(path, year):
    rows, printed = [], []
    for line in pdf_lines(path):
        match = re.match(r"^\s*(\d{1,5}|SKUPAJ)\s+", line)
        if not match:
            continue
        code = match[1]
        values = [si(n) for n in NUMBER.findall(line[match.end():])]
        assert len(values) in (6, 7, 8)
        assert values[0] + values[2] == values[4]
        assert values[1] + values[3] == values[5]
        closing = values[4] - values[5]
        if code != "SKUPAJ":
            assert values[6:] == ([abs(closing)] if closing else [])
        printed.append((code, values))
        if code.isdigit() and len(code) > 3:
            rows.append((year, code, "", values[0] - values[1], values[2], values[3], closing))
    return rows, printed


def read_named_pdf(path, year, size):
    rows = []
    pending = None
    in_table = False
    for line in pdf_lines(path):
        if re.match(r"^\s*Konto\s+Naziv\b", line):
            in_table = True
            continue
        if "Izpisano:" in line:
            in_table = False
        match = re.match(r"^\s*(\d{2,6}|SKUPAJ)\s+", line)
        if match:
            pending = None
            code = match[1]
            amounts = list(NUMBER.finditer(line[match.end():]))
            assert len(amounts) == 6
            if len(code) not in size or not code.isdigit():
                continue
            rest = line[match.end():]
            name = " ".join(rest[:amounts[0].start()].split())
            v = [si(m[0]) for m in amounts]
            pending = [year, code, name, v[0] - v[1], v[2], v[3], v[4] - v[5]]
            rows.append(pending)
        elif in_table and pending is not None and line.strip():
            assert not NUMBER.search(line)
            pending[2] += " " + " ".join(line.split())
    return [tuple(r) for r in rows]


def read_monthly(path, year):
    wb = load_workbook(path, data_only=True)
    assert len(wb.worksheets) == 1
    rows = []
    for r in wb.active.iter_rows(min_row=5, values_only=True):
        net = sum(cents(v) for v in r[3:15])
        assert net == cents(r[15])
        assert cents(r[2]) + net == cents(r[16])
        rows.append((year, r[0], r[1], cents(r[2]), net, cents(r[16])))
    wb.close()
    return rows


def read_quirky(path):
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb:
        year = int(ws.title.split()[1])
        for r in ws.iter_rows(values_only=True):
            match = re.match(r"^(\d{4,5}) - (.*)$", str(r[0]))
            if match:
                assert all(isinstance(v, str) for v in r[1:])
                rows.append((year, match[1], match[2], *(si(v) for v in r[1:])))
    wb.close()
    return rows


def read_tab(path, year):
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as f:
        lines = list(csv.reader(f, delimiter="\t"))
    assert lines[-1][0].startswith("Izpis pripravil:")
    return [(year, r[0], r[1], *(si(v) for v in r[2:])) for r in lines[2:-1]]


@pytest.fixture(scope="module", params=[("trading", 1), ("services", 7)])
def rendered(request, tmp_path_factory):
    company = generate(*request.param)
    path = tmp_path_factory.mktemp("exports")
    write_all(company, path)
    return company, path, json.loads((path / "manifest.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("fmt", FORMATS)
def test_round_trip_and_rollup(rendered, fmt):
    company, out, manifest = rendered
    files = [(name, entry) for name, entry in manifest.items() if entry["format"] == fmt]
    assert len(files) == (1 if fmt == "quirky_xlsx" else len(company.years))
    for name, entry in files:
        path = out / name
        expected = truth(out / entry["truth_file"])
        year = entry["year"]
        if fmt == "minimax_pdf":
            actual, _ = read_minimax(path, year)
        elif fmt == "analitika_pdf":
            actual = read_named_pdf(path, year, (4, 5))
        elif fmt == "furs_pdf":
            actual = read_named_pdf(path, year, (6,))
        elif fmt == "monthly_xlsx":
            actual = read_monthly(path, year)
        elif fmt == "quirky_xlsx":
            actual = read_quirky(path)
        else:
            actual = read_tab(path, year)
        target = expected
        if fmt == "monthly_xlsx":
            target = [(y, k, n, o, d - c, cl) for y, k, n, o, d, c, cl in expected]
        elif fmt == "furs_pdf":
            target = [r for r in expected if r[-1] != 0]
            hidden = [r for r in expected if not r[-1]]
            assert hidden and sum(r[3] for r in target) != 0
            assert all(r[1] in " ".join(entry["quirks"]) for r in hidden)
        assert actual == target
        rolled = defaultdict(lambda: [0, 0, 0, 0])
        assert len({(r[0], r[1]) for r in expected}) == len(expected)
        for y, k, _, o, d, c, cl in expected:
            assert o + d - c == cl and d >= 0 and c >= 0
            for i, v in enumerate((o, d, c, cl)):
                rolled[y, k[:3]][i] += v
        years = year if isinstance(year, list) else [year]
        tidy = {(y.year, r.konto): [getattr(r, k) for k in FIELDS]
                for y in company.years if y.year in years for r in y.rows}
        assert dict(rolled) == tidy


def test_minimax_subtotals_and_prefix_trap(rendered):
    _, out, manifest = rendered
    for name, entry in manifest.items():
        if entry["format"] != "minimax_pdf":
            continue
        leaves, printed = read_minimax(out / name, entry["year"])
        codes = {r[1] for r in leaves}
        assert {"1200", "12001"} <= codes
        for i, (code, values) in enumerate(printed):
            if not code.isdigit() or len(code) > 3:
                continue
            children = [(j, v) for j, (k, v) in enumerate(printed)
                        if len(k) > 3 and k.isdigit() and k[:len(code)] == code]
            assert children and all(j > i for j, _ in children)
            assert values[:6] == [sum(v[n] for _, v in children) for n in range(6)]
            assert values[4] - values[5] == sum(r[-1] for r in leaves if r[1][:len(code)] == code)
        grand = printed[-1]
        assert grand[0] == "SKUPAJ" and len(grand[1]) == 8
        assert grand[1][6] == grand[1][7]
        assert grand[1][6] == sum(max(v[4] - v[5], 0) for k, v in printed if len(k) == 1)


def test_manifest_inventory(rendered):
    _, out, manifest = rendered
    exports = {p.name for p in out.iterdir() if p.suffix in (".pdf", ".xlsx", ".csv")
               and not p.name.endswith(".truth.csv") and p.name != "statements_filed.csv"}
    assert set(manifest) == exports
    for name, entry in manifest.items():
        assert (out / name).is_file() and (out / entry["truth_file"]).is_file()
        assert entry["year"] and entry["quirks"]
        if entry["format"] in FORMATS:
            assert entry["imitates"].endswith(".pdf")


def workbook_values(path):
    wb = load_workbook(path)
    values = [(ws.title, list(ws.values)) for ws in wb]
    wb.close()
    return values


def test_deterministic_and_legacy_compatible(rendered, tmp_path):
    company, out, manifest = rendered
    write_all(company, tmp_path)
    assert (out / "manifest.json").read_bytes() == (tmp_path / "manifest.json").read_bytes()
    for name, entry in manifest.items():
        assert (out / entry["truth_file"]).read_bytes() == (tmp_path / entry["truth_file"]).read_bytes()
        if name.endswith(".xlsx"):
            assert workbook_values(out / name) == workbook_values(tmp_path / name)
        else:
            assert (out / name).read_bytes() == (tmp_path / name).read_bytes()
    baseline = tmp_path / "baseline.csv"
    write_tidy(company, baseline)
    assert baseline.read_bytes() == (out / "tb_tidy.csv").read_bytes()
    write_statements(company, baseline, AopTree.load())
    assert baseline.read_bytes() == (out / "statements_filed.csv").read_bytes()
    for y in company.years:
        write_export_csv(company, y, baseline)
        assert baseline.read_bytes() == (out / f"tb_{y.year}_export.csv").read_bytes()
        legacy = tmp_path / "baseline.xlsx"
        write_erp_xlsx(company, y, legacy)
        assert workbook_values(legacy) == workbook_values(out / f"tb_{y.year}_erp.xlsx")
