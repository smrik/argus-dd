"""Messy renderers modelled on real exports in data/tb_web/.

Known gap: reportlab text extracts cleanly; real exports scramble under
pdftotext -layout. That extraction failure is not reproduced here.
Known gap: synthetic leaves show opening on one side only; real Minimax leaves can
carry both (Litija 9330: 5.114.069,08 Breme and Dobro).
Reference statements and README are not TB exports; the manifest inventories TBs.
"""

import csv
import json
import random
import textwrap
from dataclasses import replace
from decimal import Decimal
from itertools import groupby

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from argus.layouts import eur, eur_si
from argus.tb import TbRow

FIELDS = ("opening", "debit", "credit", "closing")
SOURCES = {
    "minimax_pdf": "litija_bruto_bilanca.pdf",
    "analitika_pdf": "slovenijavgibanju_bruto_bilanca_2020.pdf",
    "furs_pdf": "kocevje_priloga1_bruto_bilanca_2018.pdf",
    "monthly_xlsx": "kamnik_velika_planina_pregled.pdf",
    "quirky_xlsx": "litija_bruto_bilanca.pdf",
    "tab_csv": "slovenijavgibanju_bruto_bilanca_2020.pdf",
    "numeric_konto_xlsx": None,
}


def ascii_text(text):
    return text.translate(str.maketrans("čšžČŠŽ", "cszCSZ"))


def parts(amount, weights):
    values = [round(Decimal(amount) * w / sum(weights)) for w in weights]
    main = max(range(len(weights)), key=weights.__getitem__)
    values[main] += amount - sum(values)
    return values


def analytic_rows(rows, rng, six_digits=False):
    """Split leaves like data/tb_web/litija_bruto_bilanca.pdf, including a prefix trap."""
    out = []
    for i, r in enumerate(rows):
        count = rng.choice((2, 3)) if r.konto == "120" or i % 3 == 0 else 1
        weights = [rng.randint(1, 9) for _ in range(count)]
        opening = parts(r.opening, weights)
        delta = parts(r.debit - r.credit, weights)
        activity = parts(min(r.debit, r.credit), weights)
        for j in range(count):
            # Four/five-digit siblings are leaves, even when one prefixes another.
            suffix = ("0", "01", "2")[j]
            konto = r.konto + (f"{j:03d}" if six_digits else suffix)
            out.append(TbRow(konto, r.name, opening[j], max(delta[j], 0) + activity[j],
                             max(-delta[j], 0) + activity[j], opening[j] + delta[j]))
    return sorted(out, key=lambda r: r.konto)


def total(konto, rows, name=""):
    return TbRow(konto, name, *(sum(getattr(r, f) for r in rows) for f in FIELDS))


def sides(value):
    return max(value, 0), max(-value, 0)


def gross(rows, minimax=False):
    opening = [sum(sides(r.opening)[i] for r in rows) for i in (0, 1)]
    debit, credit = [sum(getattr(r, f) for r in rows) for f in ("debit", "credit")]
    closing = [sum(sides(r.closing)[i] for r in rows) for i in (0, 1)]
    values = [*opening, debit, credit]
    if minimax:
        values += [opening[0] + debit, opening[1] + credit]
        closing = sides(sum(r.closing for r in rows))
    values += list(closing)
    return [eur_si(Decimal(v)) if not minimax or i < 6 or v else ""
            for i, v in enumerate(values)]


def write_truth(path, years):
    truth = path.with_name(path.name + ".truth.csv")
    with truth.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "konto", "name", *FIELDS])
        for year, rows in years:
            for r in rows:
                w.writerow([year, r.konto, r.name, *(eur(getattr(r, k)) for k in FIELDS)])
    return truth.name


def pdf(path, company, year, fmt, lines):
    """Common page furniture for the three data/tb_web/ PDF imitations."""
    width, height = landscape(A4)
    pages = []
    while lines:
        end = min(45, len(lines))
        while end < len(lines) and not lines[end][0]:
            end -= 1  # Keep wrapped names with their account across page breaks.
        pages.append(lines[:end])
        lines = lines[end:]
    c = canvas.Canvas(str(path), pagesize=(width, height), invariant=1)
    named = fmt != "minimax_pdf"
    ends = [410 + 64 * i for i in range(6)] if named else [198 + 87 * i for i in range(8)]
    for page, chunk in enumerate(pages, 1):
        c.setFont("Helvetica", 9)
        if fmt == "minimax_pdf":
            c.drawRightString(width - 24, height - 22, ascii_text(company.name))
            headers = ["Bruto bilanca", f"Temeljnice: 01.01.{year} - 31.12.{year}", "Konti: 0 - 9999999"]
            labels = ["ZACETNO", "PROMET", "SKUPAJ", "KONCNO"]
        elif fmt == "analitika_pdf":
            headers = ["DENAR | OBRACUNI", "Bruto bilanca - Analitika 1", ascii_text(company.name),
                       "Slovenska cesta 1, 1000 Ljubljana", f"Uporabnik: racunovodja   Izpis: 15.02.{year + 1} 09:00"]
            for i, text in enumerate([f"Datum obdobja knjizenja od 01.01.{year} do 31.12.{year}",
                                      "Konto 1 Vsi", "Subjekt Vsi", "Vrste dokumentov 400, 410, 420, 430"]):
                c.drawString(470, height - 22 - i * 12, text)
            labels = ["Otvoritev", "Promet", "Saldo"]
        else:
            headers = [ascii_text(company.name) + " - Slovenska cesta 1, Ljubljana",
                       "Bruto bilanca (predizdelana)", f"Bil. obdobje od: 1.{year} Bil. obdobje do: 12.{year}",
                       "Analiticna konta z zbiri na 3", "Kontov, ki imajo saldo nic ne prikazuj",
                       "Izpis v: V domaci valuti"]
            labels = ["Otvoritev", "Promet", "Saldo"]
        for i, text in enumerate(headers):
            c.drawString(24, height - 22 - 12 * i, text)
        c.setFont("Helvetica", 7)
        c.drawString(24, 492, "Konto" if named else "KONTO")
        if named:
            c.drawString(82, 492, "Naziv")
        for i, label in enumerate(labels):
            c.drawCentredString((ends[2 * i] + ends[2 * i + 1]) / 2, 505, label)
            for j, side in enumerate(("Breme", "Dobro") if not named else
                                     (("DB", "KR") if fmt == "analitika_pdf" else ("Debet", "Kredit"))):
                c.drawRightString(ends[2 * i + j], 492, side)
        for i, (konto, name, amounts) in enumerate(chunk):
            y = 477 - 9.5 * i
            c.drawString(24, y, konto)
            if named:
                c.drawString(82, y, name)
            for x, amount in zip(ends, amounts):
                c.drawRightString(x, y, amount)
        c.setFont("Helvetica", 8)
        c.drawString(24, 24, "www.minimax.si" if not named else "Bruto bilanca")
        c.drawCentredString(width / 2, 24, f"Izpisano: 15.02.{year + 1} 09:00")
        c.drawRightString(width - 24, 24, f"Stran: {page}/{len(pages)}")
        c.showPage()
    c.save()


def minimax_pdf(company, y, path, rng):
    """Imitate data/tb_web/litija_bruto_bilanca.pdf (Minimax)."""
    rows = [replace(r, name="") for r in analytic_rows(y.rows, rng)]
    lines, seen = [], set()
    for r in rows:
        for size in (1, 2, 3):
            prefix = r.konto[:size]
            if prefix not in seen:
                children = [child for child in rows if child.konto[:size] == prefix]
                lines.append((prefix, "", gross(children, minimax=True)))
                seen.add(prefix)
        lines.append((r.konto, "", gross([r], minimax=True)))
    # Like Litija: grand-total KONCNO sums razred nets by side, not leaf closings.
    nets = [sum(r.closing for r in g) for _, g in groupby(rows, key=lambda r: r.konto[0])]
    grand = gross(rows, minimax=True)[:6] + [eur_si(sum(sides(n)[i] for n in nets)) for i in (0, 1)]
    lines.append(("SKUPAJ", "", grand))
    pdf(path, company, y.year, "minimax_pdf", lines)
    return rows, ["totals before children", "4/5 digit prefix siblings are leaves", "no account names", "blank zero closing sides"]


def analitika_pdf(company, y, path, rng):
    """Imitate data/tb_web/slovenijavgibanju_bruto_bilanca_2020.pdf (Analitika 1)."""
    rows = [replace(r, name=ascii_text(r.name + (" - ureditev nepremicnin in poslovnih prostorov" if i % 7 == 0 else "")))
            for i, r in enumerate(analytic_rows(y.rows, rng))]
    lines = []

    def amounts(members):
        values = gross(members)[:4]
        net = sum(r.opening + r.debit - r.credit for r in members)
        return values + [eur_si(Decimal(v)) if v else "" for v in sides(net)]

    for konto, children in groupby(rows, key=lambda r: r.konto[:3]):
        children = list(children)
        for r in children:
            names = textwrap.wrap(r.name, width=55)
            lines.append((r.konto, names[0], amounts([r])))
            lines.extend(("", name, []) for name in names[1:])
        lines.append((konto, ascii_text(children[0].name.split(" - ")[0]).upper()[:55], amounts(children)))
    lines.append(("SKUPAJ", "", amounts(rows)))
    pdf(path, company, y.year, "analitika_pdf", lines)
    return rows, ["3 digit totals after children", "wrapped names on amount-free lines", "one-sided Saldo; blank when zero"]


def furs_pdf(company, y, path, rng):
    """Imitate data/tb_web/kocevje_priloga1_bruto_bilanca_2018.pdf (FURS)."""
    rows = [replace(r, name=ascii_text(r.name)) for r in analytic_rows(y.rows, rng, six_digits=True)]
    visible = [r for r in rows if r.closing != 0]
    lines = []
    for group, members in groupby(visible, key=lambda r: r.konto[:2]):
        members = list(members)
        for konto, children in groupby(members, key=lambda r: r.konto[:3]):
            children = list(children)
            for r in children:
                names = textwrap.wrap(r.name, width=55)
                lines.append((r.konto, names[0], gross([r])))
                lines.extend(("", name, []) for name in names[1:])
            lines.append((konto, "", gross(children)))
        lines.append((group, f"SKUPINA {group}", gross(members)))
    lines.append(("SKUPAJ", "", gross(visible)))
    pdf(path, company, y.year, "furs_pdf", lines)
    hidden = ", ".join(r.konto for r in rows if not r.closing)
    return rows, ["6 digit leaves", "totals after children", "zero closing rows hidden: " + hidden,
                  "truth includes hidden rows; printed opening need not balance"]


def monthly_xlsx(company, y, path, rng):
    """Imitate konto/month tables in data/tb_web/kamnik_velika_planina_pregled.pdf."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pregled po mesecih"
    ws.append([company.name])
    ws.append([f"Bruto bilanca {y.year} - promet po mesecih"])
    ws.append([])
    ws.append(["Konto", "Naziv", "Otvoritev", "Jan", "Feb", "Mar", "Apr", "Maj", "Jun",
               "Jul", "Avg", "Sep", "Okt", "Nov", "Dec", "Skupaj promet", "Saldo"])
    for r in y.rows:
        debit = parts(r.debit, [rng.randint(1, 20) for _ in range(12)])
        credit = parts(r.credit, [rng.randint(1, 20) for _ in range(12)])
        ws.append([r.konto, r.name, r.opening / 100, *[(d - c) / 100 for d, c in zip(debit, credit)],
                   (r.debit - r.credit) / 100, r.closing / 100])
    ws.column_dimensions["B"].width = 65
    for row in ws.iter_rows(min_row=5, min_col=3):
        for cell in row:
            cell.number_format = "#,##0.00"
    wb.save(path)
    return y.rows, ["months as columns", "monthly net movement", "gross debit/credit not recoverable; only net"]


def numeric_konto_xlsx(company, y, path, rng):
    """Imitate a generic Excel export that loses leading zeros in numeric konti."""
    rows = analytic_rows(y.rows, rng)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bruto bilanca"
    ws.append([company.name])
    ws.append([f"Bruto bilanca za leto {y.year}"])
    ws.append(["Konto", "Naziv", "Otvoritev", "Promet D", "Promet K", "Saldo"])
    for r in rows:
        ws.append([int(r.konto), r.name, *(Decimal(getattr(r, f)) / 100 for f in FIELDS)])
    wb.save(path)
    return rows, ["generic Excel export", "numeric konto cells lose leading zeros", "numeric amounts"]


def quirky_xlsx(company, path):
    """Workbook/paste variant of data/tb_web/litija_bruto_bilanca.pdf."""
    wb = Workbook()
    wb.remove(wb.active)
    truth = []
    header = ["Konto - Naziv", "Saldo začetni", "Promet D", "Promet K", "Saldo končni"]
    for y in company.years:
        rng = random.Random(f"{company.seed}-quirky_xlsx-{y.year}")
        rows = analytic_rows(y.rows, rng)
        truth.append((y.year, rows))
        ws = wb.create_sheet(f"BB {y.year}")
        ws.append([company.name])
        ws.append([f"Bruto bilanca za leto {y.year}"])
        ws.append([])
        ws.append(header)

        def append(label, r):
            amounts = []
            for f in FIELDS:
                value = getattr(r, f)
                text = eur_si(abs(value))
                if value < 0:
                    text = text + "-" if rng.choice((True, False)) else f"({text})"
                amounts.append(text)
            ws.append([label, *amounts])

        n = 0
        for klass, members in groupby(rows, key=lambda r: r.konto[0]):
            members = list(members)
            for r in members:
                if n and n % 20 == 0:
                    ws.append([])
                if n == len(rows) // 2:
                    ws.append(header)
                append(f"{r.konto} - {r.name}", r)
                n += 1
            append(f"Skupaj razred {klass}", total(klass, members))
        append("SKUPAJ", total("", rows))
        ws.column_dimensions["A"].width = 80
        for cell in ws[4]:
            cell.font = Font(bold=True)
    wb.save(path)
    return truth, ["one sheet per year", "text amounts", "trailing minus/parentheses", "code and name in one cell",
                   "spacer rows", "repeated header", "class totals"]


def tab_csv(company, y, path, rng):
    """Text variant of data/tb_web/slovenijavgibanju_bruto_bilanca_2020.pdf."""
    rows = analytic_rows(y.rows, rng, six_digits=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter="\t", quoting=csv.QUOTE_ALL)
        w.writerow([f"Bruto bilanca {y.year} - {company.name}"])
        w.writerow(["Konto", "Naziv", "Otvoritev", "Promet D", "Promet K", "Saldo"])
        for r in rows:
            w.writerow([r.konto, r.name, *(eur_si(getattr(r, k)).replace(".", "\u00a0") for k in FIELDS)])
        w.writerow(["Izpis pripravil: racunovodja"])
    return rows, ["UTF-8 BOM", "tab delimiter", "quoted fields", "NBSP thousands separator", "6 digit text codes", "footer"]


def write_exports(company, out_dir):
    """Write data/tb_web/ imitations and inventory all TB exports, including legacy layouts."""
    manifest = {}

    def record(path, fmt, source, years, quirks):
        manifest[path.name] = {"format": fmt, "imitates": source, "year": years[0][0] if len(years) == 1 else
                               [year for year, _ in years], "quirks": quirks, "truth_file": write_truth(path, years)}

    for y in company.years:
        for renderer in (minimax_pdf, analitika_pdf, furs_pdf, monthly_xlsx, numeric_konto_xlsx, tab_csv):
            fmt = renderer.__name__
            extension = fmt.rsplit("_", 1)[1]
            path = out_dir / f"tb_{y.year}_{fmt}.{extension}"
            rows, quirks = renderer(company, y, path, random.Random(f"{company.seed}-{fmt}-{y.year}"))
            record(path, fmt, SOURCES[fmt], [(y.year, rows)], quirks)
        for fmt, extension in (("erp", "xlsx"), ("export", "csv")):
            record(out_dir / f"tb_{y.year}_{fmt}.{extension}", fmt, None, [(y.year, y.rows)], ["legacy synthetic layout"])
    path = out_dir / "tb_quirky_xlsx.xlsx"
    years, quirks = quirky_xlsx(company, path)
    record(path, "quirky_xlsx", SOURCES["quirky_xlsx"], years, quirks)
    record(out_dir / "tb_tidy.csv", "tidy", None, [(y.year, y.rows) for y in company.years], ["tidy reference"])
    (out_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
