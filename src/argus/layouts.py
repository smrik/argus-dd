"""Write a synthetic company to disk: one tidy reference plus messy ERP-style exports.

The messy layouts are what wrangling will have to undo (see CONTEXT.md, Wrangling):
- *_erp.xlsx: title block, two-row merged header, debit/credit split columns,
  code and name in one cell, class subtotals and a grand total mixed into the data.
- *_export.csv: Windows-1250, ';' separated, Slovenian number format ("1.234,56").
Imitations of real software exports (Minimax, Analitika, FURS, ...) live in argus.exports.
"""

import csv
from itertools import groupby

from openpyxl import Workbook
from openpyxl.styles import Font

from argus.aop import AopTree

NON_MONETARY = {"188", "189"}  # headcount, months of operation

def eur(cents):
    return f"{cents / 100:.2f}"


def eur_si(cents):
    """Slovenian number format: thousands '.', decimals ','."""
    text = f"{abs(cents) / 100:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".")
    return f"-{text}" if cents < 0 else text


def write_tidy(company, path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "konto", "name", "opening", "debit", "credit", "closing"])
        for y in company.years:
            for r in y.rows:
                w.writerow([y.year, r.konto, r.name, eur(r.opening), eur(r.debit), eur(r.credit), eur(r.closing)])


def write_statements(company, path, tree):
    """The figures the company would have filed with AJPES (cross-check target)."""
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "statement", "aop", "label", "amount"])
        for y in company.years:
            values = {**y.statements, **y.extras}
            for aop, line in tree.lines.items():
                amount = values.get(aop, 0)
                if aop not in NON_MONETARY:
                    amount = eur(amount)
                w.writerow([y.year, line.statement, aop, line.label, amount])


def write_erp_xlsx(company, year_data, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bruto bilanca"
    bold = Font(bold=True)
    ws.append(["BRUTO BILANCA"])
    ws.append([company.name])
    ws.append([f"Obdobje: 01.01.{year_data.year} - 31.12.{year_data.year}"])
    ws.append([f"Izpisano: 15.02.{year_data.year + 1}"])
    ws.append([])
    ws.append(["Konto", "Otvoritev", None, "Promet", None, "Končno stanje", None])
    ws.append([None, "Breme", "Dobro", "Breme", "Dobro", "Breme", "Dobro"])
    for cells in ("A6:A7", "B6:C6", "D6:E6", "F6:G6"):
        ws.merge_cells(cells)
    for cell in ws[1] + ws[6] + ws[7]:
        cell.font = bold

    def amounts(rows):
        """Debit and credit columns summed separately, as ERP reports do (so the total row shows Dr = Cr)."""
        out = []
        for field in ("opening", "debit", "credit", "closing"):
            values = [getattr(r, field) for r in rows]
            if field in ("debit", "credit"):
                out.append(sum(values) / 100)
            else:
                out += [sum(v for v in values if v > 0) / 100, sum(-v for v in values if v < 0) / 100]
        return out

    for klass, rows in groupby(year_data.rows, key=lambda r: r.konto[0]):
        rows = list(rows)
        for r in rows:
            ws.append([f"{r.konto} {r.name}", *amounts([r])])
        ws.append([f"Skupaj razred {klass}", *amounts(rows)])
        for cell in ws[ws.max_row]:
            cell.font = bold
    ws.append(["SKUPAJ", *amounts(year_data.rows)])
    for cell in ws[ws.max_row]:
        cell.font = bold

    for row in ws.iter_rows(min_row=8, min_col=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 60
    wb.save(path)


def write_export_csv(company, year_data, path):
    with open(path, "w", encoding="cp1250", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([f"Bruto bilanca za leto {year_data.year} - {company.name}"])
        w.writerow(["Konto", "Naziv konta", "Otvoritveni saldo", "Promet breme", "Promet dobro", "Končni saldo"])
        for r in year_data.rows:
            w.writerow([r.konto, r.name, eur_si(r.opening), eur_si(r.debit), eur_si(r.credit), eur_si(r.closing)])
        totals = [sum(getattr(r, f) for r in year_data.rows) for f in ("opening", "debit", "credit", "closing")]
        w.writerow(["Skupaj", "", *map(eur_si, totals)])


def write_all(company, out_dir, tree=None):
    from argus.exports import write_exports

    tree = tree or AopTree.load()
    out_dir.mkdir(parents=True, exist_ok=True)
    write_tidy(company, out_dir / "tb_tidy.csv")
    write_statements(company, out_dir / "statements_filed.csv", tree)
    for y in company.years:
        write_erp_xlsx(company, y, out_dir / f"tb_{y.year}_erp.xlsx")
        write_export_csv(company, y, out_dir / f"tb_{y.year}_export.csv")
    (out_dir / "README.txt").write_text(
        f"SYNTHETIC DATA - not a real company.\n{company.name}, profile={company.profile}, seed={company.seed}\n",
        encoding="utf-8",
    )
    write_exports(company, out_dir)
